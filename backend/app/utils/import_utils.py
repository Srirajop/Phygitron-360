import fitz  # PyMuPDF
import docx
import openpyxl
import io
import json
import logging
from typing import List, Dict
from app.agents.agents import call_llm

logger = logging.getLogger(__name__)

async def extract_text_from_file(file_bytes: bytes, filename: str) -> str:
    """Extracts raw text from PDF, DOCX, or XLSX."""
    text = ""
    extension = filename.split('.')[-1].lower()

    try:
        if extension == 'pdf':
            doc = fitz.open(stream=file_bytes, filetype="pdf")
            for page in doc:
                text += page.get_text()
            doc.close()
        
        elif extension == 'docx':
            doc = docx.Document(io.BytesIO(file_bytes))
            for para in doc.paragraphs:
                text += para.text + "\n"
        
        elif extension in ['xlsx', 'xls']:
            wb = openpyxl.load_workbook(io.BytesIO(file_bytes), data_only=True)
            for sheet in wb.worksheets:
                for row in sheet.iter_rows(values_only=True):
                    text += " | ".join([str(cell) for cell in row if cell is not None]) + "\n"
        
        else:
            # Try plain text if unknown
            text = file_bytes.decode('utf-8', errors='ignore')
            
    except Exception as e:
        logger.error(f"Error extracting text from {filename}: {e}")
        raise ValueError(f"Could not read {extension.upper()} file: {str(e)}")

    return text

async def parse_questions_with_ai(text: str) -> List[Dict]:
    """Uses LLM to parse raw text into structured question objects."""
    
    system_prompt = """
    You are an expert recruitment assistant and academic instructional designer. Your task is to extract assessment questions from the provided text and return them in a strict JSON format.
    
    CRITICAL INSTRUCTION: You MUST extract EVERY SINGLE QUESTION present in the provided text. Do not summarize, do not sample, and do not stop early. If the text contains 50 questions, you must extract all 50.
    
    Supported Question Types:
    - 'mcq': Multiple choice (must have 'options' and 'correct_answer')
    - 'written': Essay/Written type where the candidate provides a text response.
    - 'coding': A programming challenge that requires the candidate to write code.
    - 'fill_in': A short-answer question with a single correct value. Provide 'correct_answer' as the accepted answer(s).
    
    Specific Rules for 'coding' questions:
    1. CRITICAL: If the question text describes a computational problem (e.g., "Given an array...", "Return the sum..."), classify it as 'coding'.
    2. NO SUMMARIZATION: For coding questions, the 'question_text' MUST include the FULL problem description, including ALL Examples, Constraints, and Notes, exactly as they appear in the source. Do NOT summarize or shorten.
    3. **Preserve Images**: If the source text contains image URLs or placeholders (like `![image](...)` or `[Image 1]`), preserve them in the description.
    4. **Markdown Formatting**: Use Markdown for the `question_text` to keep it professional and easy to read.
    2. For every 'coding' question, you MUST generate these fields accurately:
       - 'starter_code': A Python function scaffold. e.g., "def solution(input_var):\n    # Write your code here\n    pass"
       - 'test_cases': A JSON array of at least 3 objects. Each object MUST have:
         - 'input': Function arguments. IMPORTANT: Each argument MUST be on a NEW LINE. 
           Example for function sum(a, b):
           "input": "5\\n10"
         - 'expected_output': The exact expected return value as a string representation.
           Example: "15" or "[1, 2, 3]" or "True".
       - 'programming_language': "python"
    
    Example for 'coding' question:
    {
      "question_text": "### Problem Description\nGiven an array of integers nums and an integer target, return indices of the two numbers such that they add up to target.\n\n### Examples\n**Example 1:**\nInput: nums = [2,7,11,15], target = 9\nOutput: [0,1]\n\n### Constraints\n- 2 <= nums.length <= 10^4\n- -10^9 <= nums[i] <= 10^9",
      "question_type": "coding",
      "marks": 15,
      "starter_code": "def two_sum(nums, target):\n    pass",
      "test_cases": [
        {"input": "[2,7,11,15]\\n9", "expected_output": "[0,1]"}
      ],
      "programming_language": "python"
    }

    General Rules:
    1. Default to 'mcq' if 'options' exist, else 'written'.
    2. 'options' must be a list of strings.
    3. 'correct_answer' must be the text of the correct option.
    4. Provide 'marks' (default 5 for mcq, 15 for coding).
    5. Return ONLY a JSON object containing a 'questions' array.
    
    JSON Schema:
    {
      "questions": [
      {
        "question_text": "string (The COMPLETE, detailed problem description WITH Examples and Constraints. DO NOT SUMMARIZE.)",
        "question_type": "mcq|written|coding|fill_in",
        "marks": number,
        "options": ["string", "string", ...],
        "correct_answer": "string",
        "model_answer": "string (optional)",
        "starter_code": "string (optional)",
        "test_cases": [{"input": "string", "expected_output": "string"}] (optional),
        "programming_language": "string (optional)"
      }
    ]}
    """
    
    try:
        import asyncio
        chunk_size = 4000
        chunks = [text[i:i+chunk_size] for i in range(0, len(text), chunk_size)]
        
        all_valid_qs = []
        
        for idx, chunk in enumerate(chunks):
            user_prompt = f"EXTRACT ALL QUESTIONS FROM THIS TEXT AND RETURN A JSON OBJECT WITH A 'questions' ARRAY. ENFORCE ACCURACY ON TEST CASES:\n\n{chunk}"
            
            max_retries = 3
            for attempt in range(max_retries):
                try:
                    data = call_llm(system_prompt, user_prompt, max_tokens=1800)
                    
                    extracted = []
                    if isinstance(data, list):
                        extracted = data
                    elif isinstance(data, dict):
                        for key in ['questions', 'results', 'data', 'items']:
                            if key in data and isinstance(data[key], list):
                                extracted = data[key]
                                break
                        if not extracted and len(data.keys()) == 1:
                            val = list(data.values())[0]
                            if isinstance(val, list):
                                extracted = val
                        if not extracted:
                            extracted = [data]
                    
                    if not extracted or not isinstance(extracted, list):
                        logger.warning(f"AI could not find structured questions in chunk {idx+1}.")
                        break # Break retry, go to next chunk
                        
                    for q in extracted:
                        if isinstance(q, dict) and 'question_text' in q:
                            # Ensure type is valid
                            q_type = q.get('question_type', '').lower()
                            if q_type not in ['mcq', 'written', 'coding', 'fill_in']:
                                q['question_type'] = 'mcq' if 'options' in q else 'written'
                            else:
                                q['question_type'] = q_type
                            
                            # Heuristic for misclassified coding questions
                            text_lower = q['question_text'].lower()
                            coding_keywords = [
                                'write a function', 'write a program', 'implement a', 'function named',
                                'given an array', 'return indices', 'return the sum', 'time complexity',
                                'example 1:', 'input:', 'output:', 'constraints:', 'explanation:'
                            ]
                            if q['question_type'] == 'written' and any(word in text_lower for word in coding_keywords):
                                 q['question_type'] = 'coding'
            
                            # Coding specific defaults and sanitization
                            if q['question_type'] == 'coding':
                                if not q.get('programming_language'):
                                    q['programming_language'] = 'python'
                                
                                # Sanitize test cases
                                tcs = q.get('test_cases', [])
                                if not isinstance(tcs, list): tcs = []
                                
                                sanitized_tcs = []
                                for tc in tcs:
                                    if not isinstance(tc, dict): continue
                                    inp = str(tc.get("input", ""))
                                    exp = str(tc.get("expected_output", ""))
                                    
                                    sanitized_tcs.append({"input": inp, "expected_output": exp})
                                
                                q['test_cases'] = sanitized_tcs
            
                                if not q.get('starter_code'):
                                    import re
                                    match = re.search(r"function (?:named |called )?['\"]?([a-zA-Z0-9_]+)['\"]?", q['question_text'])
                                    fn_name = match.group(1) if match else "solution"
                                    q['starter_code'] = f"def {fn_name}(n):\n    # TODO: Implement\n    pass"
                            
                            all_valid_qs.append(q)
                            
                    break # success, break retry loop
                except Exception as e:
                    err_str = str(e).lower()
                    if 'rate limit' in err_str or 'rate_limit_exceeded' in err_str or '429' in err_str or '413' in err_str or 'too large' in err_str or 'tokens per minute' in err_str:
                        if attempt < max_retries - 1:
                            logger.warning(f"Rate limit or size limit hit on chunk {idx+1}/{len(chunks)}, sleeping 60s... (Attempt {attempt+1}/{max_retries})")
                            await asyncio.sleep(60)
                        else:
                            logger.error(f"Rate limit failed on chunk {idx+1} after max retries: {e}")
                            break
                    else:
                        logger.error(f"AI Parsing failed on chunk {idx+1}: {e}")
                        break
                        
            # Sleep between chunks to smooth out TPM usage
            if idx < len(chunks) - 1:
                await asyncio.sleep(2)
                
        if not all_valid_qs:
            raise ValueError("No valid questions found after AI parsing.")
            
        return all_valid_qs

    except Exception as e:
        logger.error(f"AI Parsing failed completely: {e}")
        raise ValueError(f"AI Parsing Error: {str(e)}")
